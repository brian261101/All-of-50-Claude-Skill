# -*- coding: utf-8 -*-
"""
Điền phần tóm tắt (mục II) và kết luận (mục III) vào biên bản do ghi_bien_ban_hop.py tạo ra,
thay cho các dòng chấm để trống. Phần còn lại của biên bản giữ nguyên.

Dùng:
    python dien_bien_ban.py "Biên bản - ... .docx" noi_dung.json [-o "file mới.docx"]

noi_dung.json (UTF-8), mọi khóa đều tùy chọn:
{
  "tom_tat":  ["Đoạn 1 ...", "**Phòng Kỹ thuật** báo cáo ..."],
  "ket_luan": ["1. Giao Phòng Kỹ thuật ... hoàn thành trước ngày 30/10/2026.", "2. ..."],
  "so": "15/BB-KT",
  "dia_danh": "Tây Ninh"
}
Chữ đặt trong **...** được in đậm. Không có -o thì ghi đè lên chính file biên bản.
"""
import argparse
import json
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

SUMMARY_LABEL = "Tóm tắt ý kiến phát biểu"
CONCLUSION_HEADING = "III. KẾT LUẬN"


def blank_lines_after(paragraphs, starts):
    """The dotted lines (a lone tab with a dot leader) right after the paragraph beginning with `starts`.
    None when that paragraph is missing."""
    for i, p in enumerate(paragraphs):
        if p.text.strip().startswith(starts):
            out = []
            for q in paragraphs[i + 1:]:
                if q.text != "\t":
                    break
                out.append(q)
            return out
    return None


def add_runs(p, text):
    for k, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            run = p.add_run(part)
            if k % 2:
                run.bold = True


def as_items(value):
    if not value:
        return []
    if isinstance(value, str):
        value = value.splitlines()
    return [str(v).strip() for v in value if str(v).strip()]


def fill(paragraphs, starts, items, label):
    if not items:
        return 0
    slot = blank_lines_after(paragraphs, starts)
    if slot is None:
        raise SystemExit(f"LỖI: không thấy dòng “{starts}…” trong biên bản. File này có phải do "
                         f"ghi_bien_ban_hop.py tạo ra không?")
    if not slot:
        raise SystemExit(f"LỖI: {label} không còn dòng chấm trống, có lẽ đã được điền rồi. "
                         f"Sửa trực tiếp trong Word, hoặc xuất lại biên bản từ đầu.")
    for text in items:
        p = slot[0].insert_paragraph_before()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Cm(1)
        p.paragraph_format.space_after = Pt(6)
        add_runs(p, text)
    for q in slot:
        q._p.getparent().remove(q._p)
    return len(items)


def header_paragraphs(doc):
    if not doc.tables:
        return []
    return [p for row in doc.tables[0].rows for cell in row.cells for p in cell.paragraphs]


def set_text(p, text):
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def main():
    ap = argparse.ArgumentParser(description="Điền tóm tắt và kết luận vào biên bản họp (.docx).")
    ap.add_argument("docx")
    ap.add_argument("json", help="file noi_dung.json (UTF-8)")
    ap.add_argument("-o", "--out", help="lưu ra file khác thay vì ghi đè")
    args = ap.parse_args()
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    with open(args.json, encoding="utf-8-sig") as f:
        data = json.load(f)
    doc = Document(args.docx)

    n_sum = fill(doc.paragraphs, SUMMARY_LABEL, as_items(data.get("tom_tat")), "Mục II")
    n_con = fill(doc.paragraphs, CONCLUSION_HEADING, as_items(data.get("ket_luan")), "Mục III")

    for p in header_paragraphs(doc):
        if data.get("so") and p.text.startswith("Số:") and p.runs:
            set_text(p, f"Số: {data['so'].strip()}")
        if data.get("dia_danh") and p.text.startswith("……,") and p.runs:
            set_text(p, data["dia_danh"].strip() + p.text[len("……"):])

    out = args.out or args.docx
    try:
        doc.save(out)
    except PermissionError:
        raise SystemExit(f"LỖI: không ghi được {out}. File đang mở trong Word? Đóng Word rồi chạy lại.")
    print(f"Đã điền {n_sum} đoạn tóm tắt, {n_con} ý kết luận.")
    print(f"BIÊN BẢN: {out}")


if __name__ == "__main__":
    main()
