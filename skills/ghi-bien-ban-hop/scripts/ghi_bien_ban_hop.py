# -*- coding: utf-8 -*-
"""
Ghi biên bản họp
================
Ghi âm cuộc họp qua micro, chuyển giọng nói tiếng Việt thành chữ ngay trong lúc họp
(mô hình Whisper large-v3-turbo chạy trên máy, không gửi âm thanh ra ngoài) và xuất
biên bản cuộc họp ra file Word (.docx) theo thể thức văn bản hành chính.

Chạy:
    pythonw ghi_bien_ban_hop.py                 mở cửa sổ ghi âm (hoặc bấm đúp "Ghi bien ban hop.bat")
    python  ghi_bien_ban_hop.py --setup         tải mô hình nhận dạng giọng nói nếu chưa có
    python  ghi_bien_ban_hop.py --file X.m4a    chuyển file ghi âm có sẵn thành biên bản, không mở cửa sổ
                                                (xem thêm: --help)
Cần:   pip install -r requirements.txt
Mô hình khác (máy yếu): đặt biến môi trường GHI_BIEN_BAN_MODEL=small
"""
import ctypes
import datetime as dt
import json
import os
import queue
import re
import sys
import threading
import time
import traceback
import wave

APP_DIR = os.path.join(os.environ.get("APPDATA") or os.path.expanduser("~"), "GhiBienBanHop")
os.makedirs(APP_DIR, exist_ok=True)
SETTINGS_PATH = os.path.join(APP_DIR, "settings.json")
LOG_PATH = os.path.join(APP_DIR, "log.txt")

# pythonw.exe has no console: libraries that print progress would crash without this.
if sys.stdout is None or sys.stderr is None:
    _log_stream = open(LOG_PATH, "a", encoding="utf-8", buffering=1)
    sys.stdout = sys.stdout or _log_stream
    sys.stderr = sys.stderr or _log_stream
os.environ.setdefault("HF_HUB_DISABLE_SYMLINKS_WARNING", "1")

import numpy as np
import sounddevice as sd
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SR = 16000                       # Whisper works at 16 kHz mono
MODEL_NAME = os.environ.get("GHI_BIEN_BAN_MODEL") or "large-v3-turbo"
CHUNK_MIN_S, CHUNK_MAX_S = 18, 28  # live audio is cut at the quietest point in this range
CPU_THREADS = max(4, min(12, (os.cpu_count() or 8) // 2))
NOTE_PREFIX = "GHI CHÚ:"
TS_RE = re.compile(r"^\[(\d{2}:\d{2}:\d{2})\]\s?(.*)$")
FIELDS = ("don_vi", "ten", "dia_diem", "chu_tri", "thu_ky")
TXT_NAME = "Nội dung ghi âm (chuyển chữ).txt"

# Phrases Whisper invents on silence or noise (learned from Vietnamese YouTube captions).
JUNK = ("đăng ký kênh", "đăng kí kênh", "subscribe", "ghiền mì gõ", "like và share",
        "theo dõi kênh", "cảm ơn các bạn đã theo dõi", "hẹn gặp lại các bạn trong",
        "video tiếp theo", "chúc các bạn xem video")


def log(msg):
    try:
        with open(LOG_PATH, "a", encoding="utf-8") as f:
            f.write(f"{dt.datetime.now():%Y-%m-%d %H:%M:%S} {msg}\n")
    except OSError:
        pass


def hms(sec):
    sec = int(max(0, sec))
    return f"{sec // 3600:02d}:{sec % 3600 // 60:02d}:{sec % 60:02d}"


def parse_hms(s):
    h, m, x = map(int, s.split(":"))
    return h * 3600 + m * 60 + x


def safe_name(s):
    s = re.sub(r'[<>:"/\\|?*\x00-\x1f]', " ", s)
    s = re.sub(r"\s+", " ", s).strip(" .")
    return s[:90] or "Cuoc hop"


def split_members(text):
    return [x.strip(" -•\t") for x in re.split(r"[\n;]", text) if x.strip(" -•\t")]


def documents_dir():
    try:
        buf = ctypes.create_unicode_buffer(260)
        if ctypes.windll.shell32.SHGetFolderPathW(None, 5, None, 0, buf) == 0 and buf.value:
            return buf.value  # follows OneDrive redirection of "Documents"
    except Exception:
        pass
    return os.path.join(os.path.expanduser("~"), "Documents")


def load_settings():
    try:
        with open(SETTINGS_PATH, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return {}


def default_out_root(settings):
    return settings.get("out_root") or os.path.join(documents_dir(), "Biên bản cuộc họp")


def build_prompt(ten, chu_tri, thu_ky, members):
    """Topic and names given to Whisper up front, so it spells them right."""
    text = "Biên bản cuộc họp. " + ". ".join(b for b in (ten, chu_tri, thu_ky) if b)
    return text + (f". Thành phần: {', '.join(members)}." if members else "")


def make_session_dir(out_root, label):
    folder = os.path.join(out_root, safe_name(f"{dt.datetime.now():%Y-%m-%d %Hh%M} - {label}"))
    os.makedirs(folder, exist_ok=True)
    return folder


def docx_name(ten, day):
    return safe_name(f"Biên bản - {ten or 'cuộc họp'} - {day:%d-%m-%Y}") + ".docx"


def parse_transcript(text):
    """'[hh:mm:ss] text' lines -> [(timestamp or None, text, is_note)]."""
    entries = []
    for raw in text.splitlines():
        raw = raw.strip()
        if not raw:
            continue
        m = TS_RE.match(raw)
        if not m:
            entries.append((None, raw, False))
            continue
        body = m.group(2).strip()
        is_note = body.upper().startswith(NOTE_PREFIX)
        if is_note:
            body = body[len(NOTE_PREFIX):].strip()
        entries.append((m.group(1), body, is_note))
    return entries


def keep_segment(seg):
    text = seg.text.strip()
    if not text:
        return False
    low = text.lower()
    if any(j in low for j in JUNK):
        return False
    return not (seg.no_speech_prob > 0.6 and seg.avg_logprob < -1.0)


# ---------------------------------------------------------------------------
# Speech to text
# ---------------------------------------------------------------------------

class Engine:
    """Loads the Whisper model once; every transcription shares it."""

    def __init__(self):
        self.model = None
        self.device = ""
        self.error = None
        self.ready = threading.Event()
        self.lock = threading.Lock()

    def load(self, emit):
        try:
            emit(("status", "Đang nạp mô hình nhận dạng giọng nói…"))
            from faster_whisper import WhisperModel, download_model
            try:
                path = download_model(MODEL_NAME, local_files_only=True)
            except Exception:
                emit(("status", "Đang tải mô hình nhận dạng giọng nói (khoảng 1,6 GB, chỉ tải lần đầu)…"))
                path = download_model(MODEL_NAME)
            self.model, self.device = self._open(WhisperModel, path)
            emit(("model", self.device))
        except Exception as e:
            self.error = f"{type(e).__name__}: {e}"
            log(traceback.format_exc())
            emit(("model_error", self.error))
        finally:
            self.ready.set()

    @staticmethod
    def _open(WhisperModel, path):
        try:
            import ctranslate2
            if ctranslate2.get_cuda_device_count() > 0:
                m = WhisperModel(path, device="cuda", compute_type="int8_float16")
                segs, _ = m.transcribe(np.zeros(SR, np.float32), language="vi", vad_filter=False)
                list(segs)  # fails here when the CUDA libraries are missing
                return m, "GPU"
        except Exception as e:
            log(f"GPU không dùng được, chuyển sang CPU: {e}")
        return WhisperModel(path, device="cpu", compute_type="int8", cpu_threads=CPU_THREADS), "CPU"

    def segments(self, audio, prompt=None):
        """Yields kept segments of one short (under 30 s) piece of 16 kHz audio."""
        with self.lock:
            segs, _ = self.model.transcribe(
                audio, language="vi", beam_size=5, vad_filter=True,
                vad_parameters={"min_silence_duration_ms": 500},
                initial_prompt=prompt or None, condition_on_previous_text=False)
            for s in segs:
                if keep_segment(s):
                    yield s


def load_audio(path):
    """Any audio or video file -> 16 kHz mono float32.

    Decoded here rather than by faster-whisper, whose decoder passes an argument that
    newer PyAV releases no longer accept."""
    import av
    parts = []
    with av.open(path) as container:
        if not container.streams.audio:
            raise ValueError("file không có âm thanh")
        resampler = av.AudioResampler(format="s16", layout="mono", rate=SR)
        for frame in container.decode(container.streams.audio[0]):
            parts.extend(f.to_ndarray().reshape(-1) for f in resampler.resample(frame))
        parts.extend(f.to_ndarray().reshape(-1) for f in resampler.resample(None))
    if not parts:
        raise ValueError("file không có âm thanh")
    return np.concatenate(parts).astype(np.float32) / 32768.0


def quiet_cut(buf):
    """Index of the quietest 0.1 s between CHUNK_MIN_S and CHUNK_MAX_S, so words are not split."""
    lo, hi = CHUNK_MIN_S * SR, min(len(buf), CHUNK_MAX_S * SR)
    win = SR // 10
    k = (hi - lo) // win
    if k < 1:
        return hi
    energy = (buf[lo:lo + k * win].reshape(k, win) ** 2).mean(axis=1)
    return lo + int(energy.argmin()) * win + win // 2


def transcribe_recording(engine, audio, prompt):
    """Yields (start_s, end_s, text) over a whole recording, cut at quiet points the same way
    the live recorder cuts, so sentences are not split at Whisper's fixed 30 s windows."""
    pos, tail = 0, ""
    while pos < len(audio):
        rest = audio[pos:]
        cut = quiet_cut(rest[:CHUNK_MAX_S * SR]) if len(rest) > CHUNK_MAX_S * SR else len(rest)
        p = (prompt[:200] + " " + tail[-200:]).strip()
        for s in engine.segments(rest[:cut], p):
            text = s.text.strip()
            tail = text
            yield pos / SR + s.start, pos / SR + s.end, text
        pos += cut


class Transcriber(threading.Thread):
    """Turns the live recording into text, one ~25 s piece at a time, while the meeting goes on."""

    def __init__(self, engine, prompt, emit):
        super().__init__(daemon=True)
        self.engine, self.prompt, self.emit = engine, prompt, emit
        self.q = queue.Queue()
        self.tail = ""
        self._pending = 0
        self._lock = threading.Lock()

    @property
    def pending_s(self):
        return self._pending / SR

    def feed(self, audio):
        with self._lock:
            self._pending += len(audio)
        self.q.put(audio)

    def finish(self):
        self.q.put(None)

    def run(self):
        parts, n, offset = [], 0, 0
        try:
            while True:
                a = self.q.get()
                if a is None:
                    break
                parts.append(a)
                n += len(a)
                if n >= CHUNK_MAX_S * SR:
                    buf = np.concatenate(parts)
                    cut = quiet_cut(buf)
                    self._chunk(buf[:cut], offset)
                    offset += cut
                    parts, n = [buf[cut:]], len(buf) - cut
            if n:
                self._chunk(np.concatenate(parts), offset)
        except Exception as e:
            log(traceback.format_exc())
            self.emit(("error", f"Lỗi khi chuyển giọng nói thành chữ: {e}\n\nFile ghi âm vẫn được lưu đầy đủ."))
        self.emit(("done", None))

    def _chunk(self, audio, offset):
        self.engine.ready.wait()
        try:
            if self.engine.model is None:
                return
            prompt = (self.prompt[:200] + " " + self.tail[-200:]).strip()
            for s in self.engine.segments(audio, prompt):
                text = s.text.strip()
                self.emit(("seg", (offset / SR + s.start, text)))
                self.tail = text
        finally:
            with self._lock:
                self._pending -= len(audio)


# ---------------------------------------------------------------------------
# Recording
# ---------------------------------------------------------------------------

class Recorder:
    """Microphone -> WAV file on disk (kept valid after every write) + live transcriber."""

    def __init__(self, device, on_audio):
        self.on_audio = on_audio
        self.q = queue.Queue()
        self.paused = False
        self.level = 0.0
        self.samples = 0
        self.wav = None
        self.stream = sd.InputStream(device=device, samplerate=SR, channels=1, dtype="float32",
                                     blocksize=SR // 10, callback=self._callback)
        self.thread = threading.Thread(target=self._pump, daemon=True)

    def _callback(self, indata, frames, time_info, status):
        a = indata[:, 0].copy()
        self.level = float(np.sqrt(np.mean(a * a)))
        if not self.paused:
            self.q.put(a)

    def _pump(self):
        while True:
            a = self.q.get()
            if a is None:
                break
            self.wav.writeframes((np.clip(a, -1, 1) * 32767).astype("<i2").tobytes())
            self.samples += len(a)
            self.on_audio(a)
        self.wav.close()

    def start(self, wav_path):
        self.wav = wave.open(wav_path, "wb")
        self.wav.setnchannels(1)
        self.wav.setsampwidth(2)
        self.wav.setframerate(SR)
        self.thread.start()
        self.stream.start()

    def stop(self):
        self.stream.stop()
        self.stream.close()
        self.q.put(None)
        self.thread.join()


# ---------------------------------------------------------------------------
# Word document (thể thức theo Nghị định 30/2020/NĐ-CP)
# ---------------------------------------------------------------------------

FONT = "Times New Roman"
TEXT_WIDTH_CM = 16.0  # A4 21 cm - left 3 cm - right 2 cm
_PPR_AFTER_BDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap",
                  "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi",
                  "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing",
                  "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
                  "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr",
                  "w:pPrChange")
C, J = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY


def vn_day(d):
    """Ngày < 10 và tháng 1, 2 viết thêm số 0 phía trước."""
    month = f"{d.month:02d}" if d.month <= 2 else str(d.month)
    return f"ngày {d.day:02d} tháng {month} năm {d.year}"


def signature_name(person):
    """'Ông Nguyễn Văn An - Giám đốc' -> 'Nguyễn Văn An' (chữ ký chỉ ghi họ tên)."""
    name = re.split(r"\s+[-–]\s+|,|\(", person or "")[0].strip()
    return re.sub(r"^(ông|bà|anh|chị|đ/c|đồng chí)\s+", "", name, flags=re.I)


def _run(p, text, size=None, bold=False, italic=False, color=None):
    r = p.add_run(text)
    if bold:
        r.bold = True
    if italic:
        r.italic = True
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor(*color)
    return r


def _para(box, text="", size=None, bold=False, italic=False, align=None, before=None, after=None,
          first=None, left=None, p=None):
    p = p if p is not None else box.add_paragraph()
    if text:
        _run(p, text, size, bold, italic)
    pf = p.paragraph_format
    if align is not None:
        p.alignment = align
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if first is not None:
        pf.first_line_indent = Cm(first)
    if left is not None:
        pf.left_indent = Cm(left)
    return p


def _rule(cell, inset_cm):
    """The short line drawn under the unit name and under the motto."""
    p = cell.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = pf.right_indent = Cm(inset_cm)
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = Pt(3)
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    p._p.get_or_add_pPr().insert_element_before(bdr, *_PPR_AFTER_BDR)


def _dots(doc, n):
    for _ in range(n):
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM - 0.1), WD_TAB_ALIGNMENT.RIGHT,
                                                  WD_TAB_LEADER.DOTS)
        p.paragraph_format.space_after = Pt(2)
        p.add_run("\t")


def _widths(table, *cms):
    table.autofit = False
    for i, w in enumerate(cms):
        table.columns[i].width = Cm(w)
        for cell in table.columns[i].cells:
            cell.width = Cm(w)


def build_docx(path, m, entries):
    """m: meeting info dict; entries: [(timestamp or None, text, is_note)] in meeting order."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2)
    sec.left_margin, sec.right_margin = Cm(3), Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(14)
    fonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(k), FONT)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    doc.core_properties.title = f"Biên bản {m['ten']}".strip()
    doc.core_properties.author = m["thu_ky"] or ""

    start, end = m["start"], m["end"]
    day = start.date() if start else dt.date.today()

    # Header: unit (left) and national motto (right)
    t = doc.add_table(rows=2, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    _widths(t, 6.3, 9.7)  # the motto line needs ~9.3 cm at 12 pt bold to stay on one line
    left, right = t.cell(0, 0), t.cell(0, 1)
    _para(left, (m["don_vi"] or "TÊN CƠ QUAN, ĐƠN VỊ").upper(), 13, bold=True, align=C, after=0,
          p=left.paragraphs[0])
    _rule(left, 1.9)
    _para(right, "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", 12, bold=True, align=C, after=0,
          p=right.paragraphs[0])
    _para(right, "Độc lập - Tự do - Hạnh phúc", 13, bold=True, align=C, after=0)
    _rule(right, 1.75)
    _para(t.cell(1, 0), "Số: ……/BB-……", 13, align=C, before=6, after=0, p=t.cell(1, 0).paragraphs[0])
    _para(t.cell(1, 1), f"……, {vn_day(day)}", 14, italic=True, align=C, before=6, after=0,
          p=t.cell(1, 1).paragraphs[0])

    _para(doc, "BIÊN BẢN", 14, bold=True, align=C, before=24, after=0)
    _para(doc, m["ten"] or "Cuộc họp về việc ……………………", 14, bold=True, align=C, after=12)

    # I. Time, place, attendees
    _para(doc, "I. THỜI GIAN, ĐỊA ĐIỂM VÀ THÀNH PHẦN", bold=True, before=6)
    if start:
        when = f"Bắt đầu lúc {start:%H} giờ {start:%M} phút, {vn_day(start)}."
    else:
        when = "Bắt đầu lúc …… giờ …… phút, ngày …… tháng …… năm ……"
    _para(doc, f"1. Thời gian: {when}", align=J, first=1)
    _para(doc, f"2. Địa điểm: {m['dia_diem'] or '……………………………………………'}.", align=J, first=1)
    _para(doc, f"3. Chủ trì: {m['chu_tri'] or '……………………………………………'}.", align=J, first=1)
    _para(doc, f"4. Thư ký: {m['thu_ky'] or '……………………………………………'}.", align=J, first=1)
    members = m["thanh_phan"]
    _para(doc, f"5. Thành phần tham dự{f' ({len(members)} người)' if members else ''}:", align=J,
          first=1)
    if members:
        for name in members:
            _para(doc, f"- {name}", align=J, left=1.5, after=2)
    else:
        _dots(doc, 2)

    # II. Content
    _para(doc, "II. NỘI DUNG CUỘC HỌP", bold=True, before=6)
    notes = [(ts, text) for ts, text, is_note in entries if is_note]
    if notes:
        _para(doc, "Các nội dung thư ký ghi nhanh trong cuộc họp:", align=J, first=1)
        for ts, text in notes:
            p = _para(doc, "", align=J, left=1.0, after=2)
            _run(p, f"- [{ts}] ", color=(0x59, 0x59, 0x59))
            _run(p, text)
    _para(doc, "Tóm tắt ý kiến phát biểu (toàn văn nội dung ghi âm tại Phụ lục kèm theo):",
          align=J, first=1, before=4)
    _dots(doc, 6)

    # III. Conclusion
    _para(doc, "III. KẾT LUẬN CỦA NGƯỜI CHỦ TRÌ", bold=True, before=6)
    _dots(doc, 5)

    if end:
        close = f"Cuộc họp kết thúc vào lúc {end:%H} giờ {end:%M} phút cùng ngày"
    else:
        close = "Cuộc họp kết thúc vào lúc …… giờ …… phút cùng ngày"
    closing = _para(doc, close + "; biên bản được đọc lại cho các thành viên dự họp cùng nghe.",
                    align=J, first=1, before=6, after=12)
    closing.paragraph_format.keep_with_next = True

    # Signatures: one row that may not break across pages, so names never sit alone on a page
    sig = doc.add_table(rows=1, cols=2)
    sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    _widths(sig, 8, 8)
    sig.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for cell, role, name in ((sig.cell(0, 0), "THƯ KÝ", m["thu_ky"]),
                             (sig.cell(0, 1), "CHỦ TRÌ", m["chu_tri"])):
        _para(cell, role, 14, bold=True, align=C, after=0, p=cell.paragraphs[0])
        _para(cell, "(Ký, ghi rõ họ tên)", 13, italic=True, align=C, after=0)
        for _ in range(4):
            _para(cell, "", after=0)
        _para(cell, signature_name(name), 14, bold=True, align=C, after=0)

    # Appendix: the transcript
    doc.add_page_break()
    _para(doc, "PHỤ LỤC", bold=True, align=C, after=0)
    _para(doc, "NỘI DUNG GHI ÂM CUỘC HỌP", bold=True, align=C, after=0)
    _para(doc, f"(Kèm theo Biên bản {m['ten'] or 'cuộc họp'} {vn_day(day)})", 13, italic=True,
          align=C, after=12)
    dur = f", thời lượng {hms(m['duration'])}" if m.get("duration") else ""
    _para(doc, f"Nội dung dưới đây được chuyển tự động từ giọng nói trong file ghi âm "
               f"“{m['audio_name']}”{dur}. Mốc [giờ:phút:giây] tính từ đầu file ghi âm. "
               f"Văn bản chuyển tự động có thể sai sót, cần nghe lại file ghi âm để đối chiếu "
               f"trước khi sử dụng.", 12, italic=True, align=J, after=10)
    if not entries:
        reason = m.get("transcript_error")
        _para(doc, "(Chưa có nội dung chuyển chữ" + (f": {reason}" if reason else "") + ".)", 12,
              italic=True)
    for ts, text, is_note in entries:
        p = _para(doc, "", align=J, after=3)
        if ts:
            _run(p, f"[{ts}] ", 12, color=(0x59, 0x59, 0x59))
        _run(p, f"{NOTE_PREFIX} {text}" if is_note else text, 12, bold=is_note)

    doc.save(path)


# ---------------------------------------------------------------------------
# Window
# ---------------------------------------------------------------------------

class App:
    def __init__(self, root):
        self.root = root
        self.events = queue.Queue()
        self.settings = load_settings()
        self.engine = Engine()
        self.state = "idle"  # idle | recording | paused | finishing | file
        self.recorder = None
        self.transcriber = None
        self.session = None
        self.close_after = False
        self.dirty = False
        self.saved_at = 0.0
        self.silence_check = None
        self.devices, self.default_in = [], None
        self._build_ui()
        self._load_devices()
        self.set_state("idle")
        threading.Thread(target=self.engine.load, args=(self.events.put,), daemon=True).start()
        root.protocol("WM_DELETE_WINDOW", self.on_close)
        self._poll()

    # -- layout --------------------------------------------------------------
    def _build_ui(self):
        root = self.root
        root.title("Ghi biên bản họp")
        root.geometry("1040x800")
        root.minsize(800, 620)
        base = ("Segoe UI", 10)
        root.option_add("*Font", base)
        style = ttk.Style()
        style.configure("Clock.TLabel", font=("Segoe UI", 20, "bold"))
        style.configure("Main.TButton", font=("Segoe UI", 10, "bold"), padding=(12, 6))
        style.configure("TButton", padding=(10, 5))
        style.configure("Hint.TLabel", foreground="#5b6170")

        outer = ttk.Frame(root, padding=12)
        outer.pack(fill="both", expand=True)
        outer.columnconfigure(0, weight=1)
        outer.rowconfigure(3, weight=1)

        info = ttk.LabelFrame(outer, text="Thông tin cuộc họp", padding=10)
        info.grid(row=0, column=0, sticky="ew")
        info.columnconfigure(1, weight=1)
        info.columnconfigure(3, weight=1)
        self.vars = {}

        def field(key, label, r, c, span=1):
            ttk.Label(info, text=label).grid(row=r, column=c, sticky="w", padx=(0 if c == 0 else 14, 6), pady=3)
            v = tk.StringVar(value=self.settings.get(key, ""))
            ttk.Entry(info, textvariable=v).grid(row=r, column=c + 1, columnspan=span, sticky="ew", pady=3)
            self.vars[key] = v

        field("don_vi", "Cơ quan, đơn vị", 0, 0, 3)
        field("ten", "Tên cuộc họp", 1, 0, 3)
        field("dia_diem", "Địa điểm", 2, 0)
        field("chu_tri", "Chủ trì", 2, 2)
        field("thu_ky", "Thư ký", 3, 0)
        ttk.Label(info, text="Micro").grid(row=3, column=2, sticky="w", padx=(14, 6))
        self.mic = ttk.Combobox(info, state="readonly")
        self.mic.grid(row=3, column=3, sticky="ew", pady=3)
        ttk.Label(info, text="Thành phần tham dự\n(mỗi người một dòng)").grid(row=4, column=0, sticky="nw", pady=3)
        self.members = tk.Text(info, height=3, wrap="word", font=base, relief="solid", borderwidth=1)
        self.members.grid(row=4, column=1, columnspan=3, sticky="ew", pady=3)
        self.members.insert("1.0", self.settings.get("thanh_phan", ""))
        where = ttk.Frame(info)
        where.grid(row=5, column=0, columnspan=4, sticky="ew", pady=(4, 0))
        self.where_lbl = ttk.Label(where, text="", style="Hint.TLabel")
        self.where_lbl.pack(side="left")
        ttk.Button(where, text="Đổi thư mục lưu…", command=self.change_root).pack(side="right")
        self._show_root()

        bar = ttk.Frame(outer)
        bar.grid(row=1, column=0, sticky="ew", pady=(12, 4))
        self.btn_start = ttk.Button(bar, text="●  Bắt đầu ghi", style="Main.TButton", command=self.start)
        self.btn_pause = ttk.Button(bar, text="Tạm dừng", command=self.toggle_pause)
        self.btn_stop = ttk.Button(bar, text="■  Kết thúc và xuất Word", style="Main.TButton", command=self.stop)
        for b in (self.btn_start, self.btn_pause, self.btn_stop):
            b.pack(side="left", padx=(0, 8))
        self.clock = ttk.Label(bar, text="00:00:00", style="Clock.TLabel")
        self.clock.pack(side="left", padx=(16, 10))
        self.level = ttk.Progressbar(bar, length=150, maximum=100)
        self.level.pack(side="left")
        self.btn_file = ttk.Button(bar, text="Chuyển file ghi âm có sẵn…", command=self.open_file)
        self.btn_file.pack(side="right")

        st = ttk.Frame(outer)
        st.grid(row=2, column=0, sticky="ew")
        self.status = ttk.Label(st, text="Đang khởi động…")
        self.status.pack(side="left")
        self.model_lbl = ttk.Label(st, text="", style="Hint.TLabel")
        self.model_lbl.pack(side="right")

        box = ttk.LabelFrame(outer, text="Nội dung ghi được (sửa trực tiếp được; bản sửa sẽ vào biên bản)", padding=6)
        box.grid(row=3, column=0, sticky="nsew", pady=(8, 0))
        box.columnconfigure(0, weight=1)
        box.rowconfigure(0, weight=1)
        self.txt = tk.Text(box, wrap="word", font=("Segoe UI", 11), undo=True, relief="flat",
                           padx=10, pady=8, spacing1=2, spacing3=4)
        sb = ttk.Scrollbar(box, command=self.txt.yview)
        self.txt.configure(yscrollcommand=sb.set)
        self.txt.grid(row=0, column=0, sticky="nsew")
        sb.grid(row=0, column=1, sticky="ns")
        self.txt.tag_configure("note", foreground="#9a3412", font=("Segoe UI", 11, "bold"))
        self.txt.tag_configure("hint", foreground="#7a7f8c")
        self.txt.insert("1.0", "Điền thông tin cuộc họp ở trên, rồi bấm “Bắt đầu ghi” khi cuộc họp bắt đầu.\n"
                               "Lời nói sẽ hiện ở đây sau mỗi khoảng 20–30 giây, kèm mốc thời gian.\n"
                               "Bấm “Kết thúc và xuất Word” khi họp xong để nhận file biên bản .docx.", "hint")

        nr = ttk.Frame(outer)
        nr.grid(row=4, column=0, sticky="ew", pady=(8, 0))
        nr.columnconfigure(1, weight=1)
        ttk.Label(nr, text="Ghi chú nhanh").grid(row=0, column=0, padx=(0, 8))
        self.note = ttk.Entry(nr)
        self.note.grid(row=0, column=1, sticky="ew")
        self.note.bind("<Return>", lambda e: self.add_note())
        self.btn_note = ttk.Button(nr, text="Thêm ghi chú (Enter)", command=self.add_note)
        self.btn_note.grid(row=0, column=2, padx=(8, 0))

        self.result = ttk.Frame(outer)
        self.result.grid(row=5, column=0, sticky="ew", pady=(10, 0))
        self.result_lbl = ttk.Label(self.result, text="")
        self.result_lbl.pack(side="left", fill="x", expand=True)
        ttk.Button(self.result, text="Mở thư mục", command=lambda: self._open(self.session["dir"])).pack(side="right")
        ttk.Button(self.result, text="Mở biên bản Word", style="Main.TButton",
                   command=lambda: self._open(self.session["docx"])).pack(side="right", padx=(0, 8))
        self.result.grid_remove()

    def _load_devices(self):
        try:
            devs = sd.query_devices()
            self.default_in = sd.default.device[0]
            api = devs[self.default_in]["hostapi"] if self.default_in is not None and self.default_in >= 0 else 0
            self.devices = [(i, d["name"]) for i, d in enumerate(devs)
                            if d["max_input_channels"] > 0 and d["hostapi"] == api]
        except Exception:
            log(traceback.format_exc())
            self.devices = []
        names = [n for _, n in self.devices]
        self.mic["values"] = names
        if not names:
            self.mic.set("Không tìm thấy micro")
            return
        saved = self.settings.get("mic")
        if saved in names:
            self.mic.current(names.index(saved))
        else:
            self.mic.current(next((k for k, (i, _) in enumerate(self.devices) if i == self.default_in), 0))

    # -- helpers -------------------------------------------------------------
    def out_root(self):
        return default_out_root(self.settings)

    def _show_root(self):
        self.where_lbl.config(text=f"Lưu ghi âm và biên bản vào: {self.out_root()}")

    def change_root(self):
        d = filedialog.askdirectory(title="Chọn thư mục lưu biên bản", initialdir=self.out_root(), mustexist=False)
        if d:
            self.settings["out_root"] = os.path.normpath(d)
            self._show_root()
            self.save_settings()

    def save_settings(self):
        data = {k: v.get().strip() for k, v in self.vars.items()}
        data["thanh_phan"] = self.members.get("1.0", "end-1c").strip()
        data["mic"] = self.mic.get()
        if self.settings.get("out_root"):
            data["out_root"] = self.settings["out_root"]
        self.settings = data
        try:
            with open(SETTINGS_PATH, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=1)
        except OSError:
            pass

    def prompt(self):
        ten, chu_tri, thu_ky = (self.vars[k].get().strip() for k in ("ten", "chu_tri", "thu_ky"))
        return build_prompt(ten, chu_tri, thu_ky, split_members(self.members.get("1.0", "end")))

    def _open(self, path):
        try:
            os.startfile(path)
        except OSError as e:
            messagebox.showerror("Không mở được", str(e))

    def new_session(self, label, start, audio_path):
        folder = make_session_dir(self.out_root(), label)
        return dict(dir=folder, start=start, end=None, duration=None, docx=None,
                    audio=audio_path or os.path.join(folder, "Ghi âm cuộc họp.wav"),
                    txt=os.path.join(folder, TXT_NAME))

    def set_state(self, state):
        self.state = state
        live = state in ("recording", "paused")
        idle = state == "idle"
        self.btn_start.config(state="normal" if idle else "disabled")
        self.btn_file.config(state="normal" if idle else "disabled")
        self.btn_pause.config(state="normal" if live else "disabled",
                              text="Tiếp tục ghi" if state == "paused" else "Tạm dừng")
        self.btn_stop.config(state="normal" if live else "disabled")
        self.note.config(state="normal" if live else "disabled")
        self.btn_note.config(state="normal" if live else "disabled")
        self.mic.config(state="readonly" if idle else "disabled")
        if not live:
            self.level["value"] = 0

    def insert_line(self, t, text, note=False):
        """Insert '[hh:mm:ss] text' in time order (notes typed live can precede slower transcription)."""
        follow = self.txt.yview()[1] >= 0.98
        line = f"[{hms(t)}] {text}"
        tags = ("note",) if note else ()
        last = int(self.txt.index("end-1c").split(".")[0])
        for ln in range(last, 0, -1):
            m = TS_RE.match(self.txt.get(f"{ln}.0", f"{ln}.end"))
            if m and parse_hms(m.group(1)) <= t:
                self.txt.insert(f"{ln}.end", "\n" + line, tags)
                break
        else:
            self.txt.insert("1.0", line + "\n", tags)
        if follow:
            self.txt.see("end")
        self.dirty = True

    def save_txt(self):
        if not self.session:
            return
        try:
            with open(self.session["txt"], "w", encoding="utf-8") as f:
                f.write(self.txt.get("1.0", "end-1c").strip() + "\n")
            self.dirty, self.saved_at = False, time.monotonic()
        except OSError:
            log(traceback.format_exc())

    # -- actions -------------------------------------------------------------
    def start(self):
        if not self.devices:
            messagebox.showerror("Không có micro", "Không tìm thấy micro nào. Hãy cắm micro rồi mở lại chương trình.")
            return
        device = self.devices[self.mic.current()][0]
        transcriber = Transcriber(self.engine, self.prompt(), self.events.put)
        try:
            recorder = Recorder(device, transcriber.feed)
        except Exception as e:
            log(traceback.format_exc())
            messagebox.showerror("Không mở được micro",
                                 f"{e}\n\nKiểm tra micro đã cắm và không bị chương trình khác giữ, và trong "
                                 "Cài đặt Windows > Quyền riêng tư và bảo mật > Micro đã bật "
                                 "“Cho phép ứng dụng máy tính truy cập micro”.")
            return
        now = dt.datetime.now()
        self.session = self.new_session(self.vars["ten"].get().strip() or "Cuộc họp", now, None)
        self.session["audio_name"] = os.path.basename(self.session["audio"])
        try:
            recorder.start(self.session["audio"])
        except Exception as e:
            log(traceback.format_exc())
            messagebox.showerror("Không ghi âm được", str(e))
            return
        self.recorder, self.transcriber = recorder, transcriber
        transcriber.start()
        self.save_settings()
        self.txt.delete("1.0", "end")
        self.result.grid_remove()
        self.silence_check = {"until": time.monotonic() + 6, "peak": 0.0}
        self.set_state("recording")

    def toggle_pause(self):
        if self.state == "recording":
            self.recorder.paused = True
            self.set_state("paused")
        elif self.state == "paused":
            self.recorder.paused = False
            self.set_state("recording")

    def stop(self):
        if self.state not in ("recording", "paused"):
            return
        self.recorder.stop()
        self.session["end"] = dt.datetime.now()
        self.session["duration"] = self.recorder.samples / SR
        self.clock.config(text=hms(self.session["duration"]))
        self.transcriber.finish()
        self.set_state("finishing")
        self.save_txt()

    def add_note(self):
        text = self.note.get().strip()
        if not text or self.state not in ("recording", "paused"):
            return
        self.insert_line(self.recorder.samples / SR, f"{NOTE_PREFIX} {text}", note=True)
        self.note.delete(0, "end")

    def open_file(self):
        if self.engine.ready.is_set() and self.engine.model is None:
            messagebox.showerror("Chưa có mô hình chuyển chữ", self.engine.error or "")
            return
        path = filedialog.askopenfilename(
            title="Chọn file ghi âm cuộc họp",
            filetypes=[("File âm thanh, video", "*.mp3 *.m4a *.wav *.aac *.ogg *.opus *.flac *.wma *.amr *.mp4 *.3gp *.webm"),
                       ("Tất cả các file", "*.*")])
        if not path:
            return
        path = os.path.normpath(path)
        label = self.vars["ten"].get().strip() or os.path.splitext(os.path.basename(path))[0]
        self.session = self.new_session(label, None, path)
        self.session["audio_name"] = os.path.basename(path)
        self.txt.delete("1.0", "end")
        self.result.grid_remove()
        self.clock.config(text="00:00:00")
        self.set_state("file")
        self.status.config(text="Đang chờ mô hình nhận dạng giọng nói…")
        threading.Thread(target=self._file_worker, args=(path, self.prompt()), daemon=True).start()

    def _file_worker(self, path, prompt):
        emit = self.events.put
        try:
            self.engine.ready.wait()
            if self.engine.model is None:
                raise RuntimeError(self.engine.error or "chưa nạp được mô hình")
            emit(("status", "Đang đọc file ghi âm…"))
            audio = load_audio(path)
            duration = len(audio) / SR
            emit(("file_info", duration))
            for start, end, text in transcribe_recording(self.engine, audio, prompt):
                emit(("seg", (start, text)))
                emit(("progress", end / duration))
        except Exception as e:
            log(traceback.format_exc())
            emit(("error", f"Không chuyển chữ được file này:\n{e}"))
        emit(("done", None))

    def export(self):
        s = self.session
        entries = parse_transcript(self.txt.get("1.0", "end-1c"))
        meta = {k: v.get().strip() for k, v in self.vars.items()}
        meta.update(thanh_phan=split_members(self.members.get("1.0", "end")), start=s["start"], end=s["end"],
                    duration=s["duration"], audio_name=s["audio_name"], transcript_error=self.engine.error)
        path = os.path.join(s["dir"], docx_name(meta["ten"], s["start"] or dt.datetime.now()))
        self.save_txt()
        try:
            build_docx(path, meta, entries)
        except Exception as e:
            log(traceback.format_exc())
            messagebox.showerror("Không lưu được biên bản Word", f"{e}\n\nNội dung chuyển chữ vẫn có trong:\n{s['txt']}")
            self.set_state("idle")
            return
        s["docx"] = path
        self.save_settings()
        self.set_state("idle")
        self.status.config(text="Đã xuất biên bản Word. Phần II, III để trống cho thư ký tóm tắt và ghi kết luận.")
        self.result_lbl.config(text=path)
        self.result.grid()
        if self.close_after:
            self.root.destroy()
            return
        self._open(path)

    # -- event loop ----------------------------------------------------------
    def handle(self, kind, data):
        if kind == "status":
            self.status.config(text=data)
        elif kind == "model":
            self.model_lbl.config(text=f"Mô hình {MODEL_NAME} · chạy bằng {data}")
            if self.state == "idle":
                self.status.config(text="Sẵn sàng. Bấm “Bắt đầu ghi” khi cuộc họp bắt đầu.")
        elif kind == "model_error":
            self.model_lbl.config(text="Không nạp được mô hình chuyển chữ")
            self.status.config(text="Vẫn ghi âm được, nhưng chưa chuyển được giọng nói thành chữ.")
            messagebox.showwarning("Không nạp được mô hình chuyển chữ",
                                   f"{data}\n\nVẫn ghi âm và xuất biên bản Word được, nhưng chưa có phần chuyển chữ. "
                                   "Lần đầu sử dụng cần có Internet để tải mô hình (khoảng 1,6 GB).")
        elif kind == "seg":
            t, text = data
            self.insert_line(t, text)
        elif kind == "file_info":
            self.session["duration"] = data
            self.clock.config(text=hms(data))
        elif kind == "progress":
            self.status.config(text=f"Đang chuyển chữ file ghi âm: {min(100, data * 100):.0f}%")
            self.level["value"] = min(100, data * 100)
        elif kind == "error":
            messagebox.showerror("Lỗi chuyển chữ", data)
        elif kind == "done":
            self.export()

    def _poll(self):
        try:
            while True:
                self.handle(*self.events.get_nowait())
        except queue.Empty:
            pass
        if self.state in ("recording", "paused"):
            rec = self.recorder
            self.clock.config(text=hms(rec.samples / SR))
            level = rec.level if self.state == "recording" else 0.0
            self.level["value"] = max(0.0, min(100.0, (20 * np.log10(max(level, 1e-6)) + 60) / 60 * 100))
            msg = "Đang ghi âm" if self.state == "recording" else "Đã tạm dừng ghi (phần tạm dừng không được ghi)"
            if not self.engine.ready.is_set():
                msg += " · đang nạp mô hình, lời nói sẽ được chuyển chữ ngay khi nạp xong"
            elif self.transcriber.pending_s > 45:
                msg += f" · phần chuyển chữ đang chậm sau khoảng {int(self.transcriber.pending_s)} giây"
            self.status.config(text=msg)
            chk = self.silence_check
            if chk and self.state == "recording":
                chk["peak"] = max(chk["peak"], level)
                if time.monotonic() > chk["until"]:
                    self.silence_check = None
                    if chk["peak"] < 1e-5:
                        messagebox.showwarning(
                            "Micro không thu được tiếng",
                            "Mấy giây đầu micro hoàn toàn im lặng. Kiểm tra: micro có bị tắt tiếng không, đã chọn "
                            "đúng micro chưa, và Cài đặt Windows > Quyền riêng tư và bảo mật > Micro đã bật "
                            "“Cho phép ứng dụng máy tính truy cập micro” chưa.\n\nChương trình vẫn đang ghi.")
        elif self.state == "finishing":
            left = int(self.transcriber.pending_s) if self.transcriber else 0
            self.status.config(text=f"Đang chuyển chữ phần cuối (còn khoảng {left} giây âm thanh), "
                                    "xong sẽ tự xuất biên bản Word…")
        if self.dirty and time.monotonic() - self.saved_at > 10:
            self.save_txt()
        self.root.after(150, self._poll)

    def on_close(self):
        if self.state in ("recording", "paused"):
            ans = messagebox.askyesnocancel(
                "Đang ghi âm",
                "Cuộc họp đang được ghi âm.\n\n"
                "Có: kết thúc ghi, xuất biên bản Word rồi đóng.\n"
                "Không: đóng ngay (file ghi âm đến lúc này vẫn được lưu).\n"
                "Hủy: quay lại.")
            if ans is None:
                return
            if ans:
                self.close_after = True
                self.stop()
                return
            self.recorder.stop()
            self.save_txt()
        elif self.state in ("finishing", "file"):
            if not messagebox.askyesno("Đang chuyển chữ",
                                       "Chương trình đang chuyển giọng nói thành chữ. Đóng ngay thì biên bản Word "
                                       "chưa được xuất (file ghi âm vẫn còn). Vẫn đóng?"):
                return
            self.save_txt()
        self.root.destroy()


# ---------------------------------------------------------------------------
# Command line (used by the launcher and by Claude)
# ---------------------------------------------------------------------------

def _parse_start(text):
    for fmt in ("%Y-%m-%d %H:%M", "%d/%m/%Y %H:%M", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(text.strip(), fmt)
        except ValueError:
            pass
    raise SystemExit(f"LỖI: không hiểu --bat-dau '{text}'. Viết dạng 2026-10-21 08:30 hoặc 21/10/2026 08:30.")


def run_cli(argv):
    import argparse
    ap = argparse.ArgumentParser(
        prog="ghi_bien_ban_hop.py",
        description="Ghi biên bản họp. Chạy không có tham số để mở cửa sổ ghi âm.")
    ap.add_argument("--setup", action="store_true",
                    help="tải mô hình nhận dạng giọng nói nếu chưa có (khoảng 1,6 GB, chỉ lần đầu)")
    ap.add_argument("--file", help="file ghi âm có sẵn (mp3, m4a, wav, mp4…) cần chuyển thành biên bản")
    ap.add_argument("--ten", help="tên cuộc họp")
    ap.add_argument("--don-vi", dest="don_vi", help="cơ quan, đơn vị (mặc định: lần dùng trước)")
    ap.add_argument("--dia-diem", dest="dia_diem", help="địa điểm họp")
    ap.add_argument("--chu-tri", dest="chu_tri", help="người chủ trì")
    ap.add_argument("--thu-ky", dest="thu_ky", help="thư ký (mặc định: lần dùng trước)")
    ap.add_argument("--thanh-phan", dest="thanh_phan", help='thành phần tham dự, ngăn cách bằng dấu ";"')
    ap.add_argument("--bat-dau", dest="bat_dau", help='giờ bắt đầu họp, ví dụ "2026-10-21 08:30"')
    ap.add_argument("--out", help="thư mục lưu (mặc định: Documents\\Biên bản cuộc họp)")
    args = ap.parse_args(argv)
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

    if args.setup:
        from faster_whisper import download_model
        try:
            path = download_model(MODEL_NAME, local_files_only=True)
        except Exception:
            print(f"Đang tải mô hình nhận dạng giọng nói {MODEL_NAME} (khoảng 1,6 GB, chỉ tải lần đầu)…", flush=True)
            path = download_model(MODEL_NAME)
        print(f"Mô hình đã sẵn sàng: {path}")
        return 0

    if not args.file:
        ap.print_help()
        return 2
    audio = os.path.abspath(args.file)
    if not os.path.isfile(audio):
        print(f"LỖI: không thấy file {audio}")
        return 1

    settings = load_settings()
    fields = {k: (getattr(args, k) or "").strip() for k in FIELDS}
    for k in ("don_vi", "thu_ky"):  # the same for every meeting of one secretary
        fields[k] = fields[k] or settings.get(k, "")
    members = split_members(args.thanh_phan or "")
    start = _parse_start(args.bat_dau) if args.bat_dau else None

    def emit(event):
        kind, data = event
        if kind == "status":
            print(data, flush=True)
        elif kind == "model":
            print(f"Đã nạp mô hình {MODEL_NAME}, chạy bằng {data}.", flush=True)
        elif kind == "model_error":
            print(f"LỖI: không nạp được mô hình: {data}", flush=True)

    engine = Engine()
    engine.load(emit)
    if engine.model is None:
        return 1

    print(f"Đang đọc file: {audio}", flush=True)
    samples = load_audio(audio)
    duration = len(samples) / SR
    print(f"Đang chuyển chữ {hms(duration)} âm thanh…", flush=True)
    prompt = build_prompt(fields["ten"], fields["chu_tri"], fields["thu_ky"], members)
    entries, shown = [], -5
    for seg_start, seg_end, text in transcribe_recording(engine, samples, prompt):
        entries.append((hms(seg_start), text, False))
        pct = int(seg_end / duration * 100)
        if pct >= shown + 5:
            print(f"  {pct}%  ({hms(seg_end)} / {hms(duration)})", flush=True)
            shown = pct

    folder = make_session_dir(args.out or default_out_root(settings),
                              fields["ten"] or os.path.splitext(os.path.basename(audio))[0])
    txt = os.path.join(folder, TXT_NAME)
    with open(txt, "w", encoding="utf-8") as f:
        f.write("\n".join(f"[{ts}] {text}" for ts, text, _ in entries) + "\n")
    meta = dict(fields, thanh_phan=members, start=start,
                end=start + dt.timedelta(seconds=duration) if start and duration else None,
                duration=duration, audio_name=os.path.basename(audio), transcript_error=None)
    docx = os.path.join(folder, docx_name(fields["ten"], start or dt.datetime.now()))
    build_docx(docx, meta, entries)
    print(f"Xong: {len(entries)} đoạn, thời lượng {hms(duration or 0)}.")
    print(f"BIÊN BẢN: {docx}")
    print(f"NỘI DUNG: {txt}")
    return 0


def main():
    if len(sys.argv) > 1:
        sys.exit(run_cli(sys.argv[1:]))
    try:
        ctypes.windll.shcore.SetProcessDpiAwareness(1)
    except Exception:
        pass
    root = tk.Tk()

    def report(exc, val, tb):
        log("".join(traceback.format_exception(exc, val, tb)))
        messagebox.showerror("Lỗi", f"{val}\n\nChi tiết lưu tại {LOG_PATH}")

    root.report_callback_exception = report
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
